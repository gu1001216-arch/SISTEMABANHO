"""
Sistema de Monitoramento de Tratamento de Peças — Pintura Eletrostática
Fluxo: Preparação (operador) -> Validação (líder) -> Banho (operador de banho)

Banco: PostgreSQL (Railway).  Em ambiente local sem DATABASE_URL, usa SQLite.
"""
import os
import io
import csv
import json
from datetime import datetime, timedelta
from functools import wraps

from flask import (
    Flask, render_template, request, redirect, url_for,
    session, jsonify, send_file, flash
)
from werkzeug.security import generate_password_hash, check_password_hash
from sqlalchemy import (
    create_engine, Column, Integer, String, DateTime, Float, Text, func
)
from sqlalchemy.orm import declarative_base, sessionmaker, scoped_session
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment

# ─────────────────────────────────────────────────────────────────────────────
# Configuração
# ─────────────────────────────────────────────────────────────────────────────
app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY', 'troque-esta-chave-em-producao')

# Railway fornece DATABASE_URL no formato postgres://; SQLAlchemy quer postgresql://
DATABASE_URL = os.environ.get('DATABASE_URL', '')
if DATABASE_URL.startswith('postgres://'):
    DATABASE_URL = DATABASE_URL.replace('postgres://', 'postgresql://', 1)
if not DATABASE_URL:
    # fallback local para desenvolvimento
    DATABASE_URL = 'sqlite:///dados_local.db'

engine = create_engine(
    DATABASE_URL,
    pool_pre_ping=True,      # reconecta se a conexão cair (essencial 24/7)
    pool_recycle=280,
)
Session = scoped_session(sessionmaker(bind=engine))
Base = declarative_base()

PROCESSOS = [
    "AÇO SEM OXIDAÇÃO",
    "AÇO COM OXIDAÇÃO",
    "ALUMÍNIO",
    "MINIMIZADO SEM OXIDAÇÃO",
    "MINIMIZADO COM OXIDAÇÃO",
    "INOX",
]

# Estados do card no fluxo
ST_PREPARANDO = 'PREPARANDO'      # operador iniciou, enchendo o cesto
ST_AGUARD_LIDER = 'AGUARDA_LIDER' # operador finalizou prep, aguarda validação
ST_NA_FILA_BANHO = 'FILA_BANHO'   # líder validou, aguarda operador de banho
ST_EM_BANHO = 'EM_BANHO'          # operador de banho iniciou
ST_CONCLUIDO = 'CONCLUIDO'        # finalizado

# ─────────────────────────────────────────────────────────────────────────────
# Modelos
# ─────────────────────────────────────────────────────────────────────────────
class Usuario(Base):
    __tablename__ = 'usuarios'
    id = Column(Integer, primary_key=True)
    login = Column(String(50), unique=True, nullable=False)
    nome = Column(String(120), nullable=False)
    senha_hash = Column(String(255), nullable=False)
    perfil = Column(String(20), nullable=False)  # admin, prep, lider, banho

    def to_dict(self):
        return {'id': self.id, 'login': self.login, 'nome': self.nome, 'perfil': self.perfil}


class ItemMestre(Base):
    """Lista mestra: OP -> código, descrição, quantidade."""
    __tablename__ = 'itens_mestre'
    id = Column(Integer, primary_key=True)
    op = Column(String(60), unique=True, nullable=False, index=True)
    codigo = Column(String(60), default='')
    descricao = Column(String(255), default='')
    quantidade = Column(Integer, default=0)

    def to_dict(self):
        return {'op': self.op, 'codigo': self.codigo,
                'descricao': self.descricao, 'quantidade': self.quantidade}


class Card(Base):
    """Card único que percorre todo o fluxo preparação -> banho."""
    __tablename__ = 'cards'
    id = Column(Integer, primary_key=True)
    estado = Column(String(20), nullable=False, index=True)

    numero_cesto = Column(String(40), nullable=False)
    processo = Column(String(60), default='')
    tipo = Column(String(20), default='Normal')   # Normal / Retrabalho

    # dados da peça / OP
    op = Column(String(60), default='')
    codigo = Column(String(60), default='')
    descricao = Column(String(255), default='')
    quantidade = Column(Integer, default=0)
    observacao = Column(Text, default='')

    # quem
    operador_prep = Column(String(120), default='')
    lider = Column(String(120), default='')
    operador_banho = Column(String(120), default='')

    # tempos de preparação
    prep_inicio = Column(DateTime)
    prep_fim = Column(DateTime)
    prep_minutos = Column(Float, default=0)

    # tempos de banho
    banho_inicio = Column(DateTime)
    banho_fim = Column(DateTime)
    banho_minutos = Column(Float, default=0)

    criado_em = Column(DateTime, default=datetime.utcnow)

    def to_dict(self):
        def fmt(dt):
            if not dt:
                return ''
            # converte de UTC para horário de Brasília (UTC-3) só na exibição
            return (dt - timedelta(hours=3)).strftime('%d/%m/%Y %H:%M:%S')
        return {
            'id': self.id,
            'estado': self.estado,
            'numero_cesto': self.numero_cesto,
            'processo': self.processo,
            'tipo': self.tipo,
            'op': self.op,
            'codigo': self.codigo,
            'descricao': self.descricao,
            'quantidade': self.quantidade,
            'observacao': self.observacao or '',
            'operador_prep': self.operador_prep,
            'lider': self.lider,
            'operador_banho': self.operador_banho,
            'prep_inicio': fmt(self.prep_inicio),
            'prep_fim': fmt(self.prep_fim),
            'prep_minutos': round(self.prep_minutos or 0, 1),
            'banho_inicio': fmt(self.banho_inicio),
            'banho_fim': fmt(self.banho_fim),
            'banho_minutos': round(self.banho_minutos or 0, 1),
            # sufixo 'Z' => o JS interpreta como UTC e converte p/ o fuso local corretamente
            'prep_inicio_iso': self.prep_inicio.isoformat() + 'Z' if self.prep_inicio else '',
            'banho_inicio_iso': self.banho_inicio.isoformat() + 'Z' if self.banho_inicio else '',
        }


# ─────────────────────────────────────────────────────────────────────────────
# Inicialização do banco + seed
# ─────────────────────────────────────────────────────────────────────────────
def init_db():
    Base.metadata.create_all(engine)
    db = Session()
    try:
        if db.query(Usuario).count() == 0:
            seed = [
                ('admin', 'Administrador', 'admin123', 'admin'),
                ('lider', 'Líder de Preparação', 'lider123', 'lider'),
                ('banho', 'Operador de Banho', 'banho123', 'banho'),
            ]
            for i in range(1, 10):
                seed.append((f'op{i}', f'Operador {i}', 'op1234', 'prep'))
            for login, nome, senha, perfil in seed:
                db.add(Usuario(login=login, nome=nome,
                               senha_hash=generate_password_hash(senha),
                               perfil=perfil))
            db.commit()
    finally:
        db.close()


# ─────────────────────────────────────────────────────────────────────────────
# Auth helpers
# ─────────────────────────────────────────────────────────────────────────────
def login_required(*perfis):
    def deco(f):
        @wraps(f)
        def wrapper(*a, **kw):
            if 'usuario' not in session:
                return redirect(url_for('login'))
            if perfis and session.get('perfil') not in perfis and session.get('perfil') != 'admin':
                return redirect(url_for('login'))
            return f(*a, **kw)
        return wrapper
    return deco


# ─────────────────────────────────────────────────────────────────────────────
# Rotas de página
# ─────────────────────────────────────────────────────────────────────────────
@app.route('/')
def index():
    return redirect(url_for('login'))


@app.route('/login', methods=['GET', 'POST'])
def login():
    erro = None
    if request.method == 'POST':
        login_u = request.form.get('usuario', '').strip()
        senha = request.form.get('senha', '')
        db = Session()
        try:
            u = db.query(Usuario).filter_by(login=login_u).first()
            if u and check_password_hash(u.senha_hash, senha):
                session['usuario'] = u.login
                session['nome'] = u.nome
                session['perfil'] = u.perfil
                # encaminha para a tela do perfil
                destino = {
                    'admin': 'dashboard',
                    'lider': 'tela_lider',
                    'banho': 'tela_banho',
                    'prep': 'tela_prep',
                }.get(u.perfil, 'login')
                return redirect(url_for(destino))
            erro = 'Usuário ou senha incorretos.'
        finally:
            db.close()
    return render_template('login.html', erro=erro)


@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login'))


@app.route('/preparacao')
@login_required('prep', 'lider')
def tela_prep():
    return render_template('prep.html', nome=session.get('nome'),
                           perfil=session.get('perfil'), processos=PROCESSOS)


@app.route('/lider')
@login_required('lider')
def tela_lider():
    return render_template('lider.html', nome=session.get('nome'),
                           perfil=session.get('perfil'), processos=PROCESSOS)


@app.route('/banho')
@login_required('banho')
def tela_banho():
    return render_template('banho.html', nome=session.get('nome'),
                           perfil=session.get('perfil'))


@app.route('/dashboard')
@login_required('admin')
def dashboard():
    db = Session()
    try:
        concluidos = db.query(Card).filter_by(estado=ST_CONCLUIDO).all()
        total = len(concluidos)
        normais = sum(1 for c in concluidos if c.tipo == 'Normal')
        retrab = sum(1 for c in concluidos if c.tipo == 'Retrabalho')
        tempos_prep = [c.prep_minutos for c in concluidos if c.prep_minutos]
        tempos_banho = [c.banho_minutos for c in concluidos if c.banho_minutos]
        media_prep = round(sum(tempos_prep) / len(tempos_prep), 1) if tempos_prep else 0
        media_banho = round(sum(tempos_banho) / len(tempos_banho), 1) if tempos_banho else 0
        por_proc = {}
        for c in concluidos:
            p = c.processo or 'Desconhecido'
            por_proc[p] = por_proc.get(p, 0) + 1
        registros = [c.to_dict() for c in
                     db.query(Card).filter_by(estado=ST_CONCLUIDO)
                     .order_by(Card.id.desc()).limit(100).all()]
        # cards em andamento (qualquer estado != concluído)
        em_andamento = db.query(Card).filter(Card.estado != ST_CONCLUIDO).count()
        return render_template(
            'dashboard.html', nome=session.get('nome'),
            total=total, normais=normais, retrabalhos=retrab,
            media_prep=media_prep, media_banho=media_banho,
            em_andamento=em_andamento,
            por_processo=por_proc, registros=registros, processos=PROCESSOS)
    finally:
        db.close()


@app.route('/admin/usuarios', methods=['GET', 'POST'])
@login_required('admin')
def admin_usuarios():
    db = Session()
    msg = None
    try:
        if request.method == 'POST':
            acao = request.form.get('acao')
            if acao == 'adicionar':
                nu = request.form.get('novo_usuario', '').strip()
                nn = request.form.get('novo_nome', '').strip()
                ns = request.form.get('nova_senha', '')
                npf = request.form.get('novo_perfil', 'prep')
                if nu and nn and ns:
                    if not db.query(Usuario).filter_by(login=nu).first():
                        db.add(Usuario(login=nu, nome=nn,
                                       senha_hash=generate_password_hash(ns),
                                       perfil=npf))
                        db.commit()
                        msg = f'Usuário {nn} adicionado.'
                    else:
                        msg = 'Esse login já existe.'
            elif acao == 'remover':
                ur = request.form.get('usuario_remover')
                u = db.query(Usuario).filter_by(login=ur).first()
                if u and u.login != 'admin':
                    db.delete(u)
                    db.commit()
                    msg = 'Usuário removido.'
            elif acao == 'senha':
                ur = request.form.get('usuario_senha')
                nova = request.form.get('senha_nova', '')
                u = db.query(Usuario).filter_by(login=ur).first()
                if u and nova:
                    u.senha_hash = generate_password_hash(nova)
                    db.commit()
                    msg = f'Senha de {u.nome} atualizada.'
        usuarios = [u.to_dict() for u in db.query(Usuario).order_by(Usuario.id).all()]
        return render_template('usuarios.html', usuarios=usuarios,
                               nome=session.get('nome'), msg=msg)
    finally:
        db.close()


@app.route('/admin/mestre', methods=['GET', 'POST'])
@login_required('admin')
def admin_mestre():
    db = Session()
    msg = None
    try:
        if request.method == 'POST':
            f = request.files.get('arquivo')
            if f and f.filename:
                nome = f.filename.lower()
                linhas = []
                try:
                    if nome.endswith('.csv'):
                        texto = io.StringIO(f.stream.read().decode('utf-8-sig'))
                        leitor = csv.reader(texto, delimiter=detectar_sep(texto))
                        linhas = list(leitor)
                    else:  # xlsx
                        from openpyxl import load_workbook
                        wb = load_workbook(f, read_only=True, data_only=True)
                        ws = wb.active
                        for row in ws.iter_rows(values_only=True):
                            linhas.append(list(row))
                    novos, atualizados = importar_mestre(db, linhas)
                    msg = f'Importado: {novos} novos, {atualizados} atualizados.'
                except Exception as e:
                    msg = f'Erro ao importar: {e}'
        total_itens = db.query(ItemMestre).count()
        amostra = [i.to_dict() for i in db.query(ItemMestre).limit(20).all()]
        return render_template('mestre.html', nome=session.get('nome'),
                               msg=msg, total_itens=total_itens, amostra=amostra)
    finally:
        db.close()


def detectar_sep(stringio):
    pos = stringio.tell()
    primeira = stringio.readline()
    stringio.seek(pos)
    return ';' if primeira.count(';') > primeira.count(',') else ','


def importar_mestre(db, linhas):
    """Espera colunas: OP, Código, Descrição, Quantidade. Ignora cabeçalho."""
    novos = atualizados = 0
    for i, row in enumerate(linhas):
        if not row or all(c is None or str(c).strip() == '' for c in row):
            continue
        # pula cabeçalho (primeira linha com texto 'op' por ex.)
        if i == 0 and str(row[0]).strip().lower() in ('op', 'ordem', 'ordem de produção'):
            continue
        op = str(row[0]).strip() if len(row) > 0 and row[0] is not None else ''
        if not op:
            continue
        codigo = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ''
        descricao = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ''
        try:
            qtd = int(float(row[3])) if len(row) > 3 and row[3] not in (None, '') else 0
        except (ValueError, TypeError):
            qtd = 0
        existente = db.query(ItemMestre).filter_by(op=op).first()
        if existente:
            existente.codigo = codigo
            existente.descricao = descricao
            existente.quantidade = qtd
            atualizados += 1
        else:
            db.add(ItemMestre(op=op, codigo=codigo, descricao=descricao, quantidade=qtd))
            novos += 1
    db.commit()
    return novos, atualizados


# ─────────────────────────────────────────────────────────────────────────────
# APIs — fluxo
# ─────────────────────────────────────────────────────────────────────────────
@app.route('/api/buscar_op/<path:op>')
@login_required('prep', 'lider')
def api_buscar_op(op):
    db = Session()
    try:
        item = db.query(ItemMestre).filter_by(op=op.strip()).first()
        if item:
            return jsonify({'encontrado': True, **item.to_dict()})
        return jsonify({'encontrado': False, 'op': op})
    finally:
        db.close()


@app.route('/api/prep/iniciar', methods=['POST'])
@login_required('prep', 'lider')
def api_prep_iniciar():
    d = request.json or {}
    numero = str(d.get('numero_cesto', '')).strip()
    if not numero:
        return jsonify({'sucesso': False, 'erro': 'Informe o número do cesto.'}), 400
    db = Session()
    try:
        # impede dois cards ativos com mesmo cesto (estados não concluídos)
        existe = db.query(Card).filter(
            Card.numero_cesto == numero,
            Card.estado != ST_CONCLUIDO
        ).first()
        if existe:
            return jsonify({'sucesso': False,
                            'erro': f'Cesto {numero} já está em um card ativo.'}), 400
        card = Card(
            estado=ST_PREPARANDO,
            numero_cesto=numero,
            operador_prep=session.get('nome', ''),
            prep_inicio=datetime.utcnow(),
        )
        db.add(card)
        db.commit()
        return jsonify({'sucesso': True, 'id': card.id})
    finally:
        db.close()


@app.route('/api/prep/finalizar', methods=['POST'])
@login_required('prep', 'lider')
def api_prep_finalizar():
    d = request.json or {}
    db = Session()
    try:
        card = db.query(Card).get(int(d.get('id', 0)))
        if not card or card.estado != ST_PREPARANDO:
            return jsonify({'sucesso': False, 'erro': 'Card não encontrado.'}), 404
        card.prep_fim = datetime.utcnow()
        card.prep_minutos = round((card.prep_fim - card.prep_inicio).total_seconds() / 60, 1)
        card.processo = d.get('processo', '')
        card.tipo = d.get('tipo', 'Normal')
        card.op = str(d.get('op', '')).strip()
        card.codigo = str(d.get('codigo', '')).strip()
        card.descricao = str(d.get('descricao', '')).strip()
        try:
            card.quantidade = int(d.get('quantidade') or 0)
        except (ValueError, TypeError):
            card.quantidade = 0
        card.observacao = d.get('observacao', '')
        card.estado = ST_AGUARD_LIDER
        db.commit()
        return jsonify({'sucesso': True})
    finally:
        db.close()


@app.route('/api/prep/ativos')
@login_required('prep', 'lider')
def api_prep_ativos():
    db = Session()
    try:
        cards = db.query(Card).filter_by(estado=ST_PREPARANDO).order_by(Card.prep_inicio).all()
        return jsonify([c.to_dict() for c in cards])
    finally:
        db.close()


@app.route('/api/lider/pendentes')
@login_required('lider')
def api_lider_pendentes():
    db = Session()
    try:
        cards = db.query(Card).filter_by(estado=ST_AGUARD_LIDER).order_by(Card.prep_fim).all()
        return jsonify([c.to_dict() for c in cards])
    finally:
        db.close()


@app.route('/api/lider/validar', methods=['POST'])
@login_required('lider')
def api_lider_validar():
    d = request.json or {}
    db = Session()
    try:
        card = db.query(Card).get(int(d.get('id', 0)))
        if not card or card.estado != ST_AGUARD_LIDER:
            return jsonify({'sucesso': False, 'erro': 'Card não encontrado.'}), 404
        # líder pode ajustar tempo de preparação e demais campos
        if d.get('prep_minutos') not in (None, ''):
            try:
                card.prep_minutos = round(float(d.get('prep_minutos')), 1)
            except (ValueError, TypeError):
                pass
        for campo in ('processo', 'tipo', 'op', 'codigo', 'descricao', 'observacao'):
            if campo in d:
                setattr(card, campo, d.get(campo))
        if 'quantidade' in d:
            try:
                card.quantidade = int(d.get('quantidade') or 0)
            except (ValueError, TypeError):
                pass
        card.lider = session.get('nome', '')
        card.estado = ST_NA_FILA_BANHO
        db.commit()
        return jsonify({'sucesso': True})
    finally:
        db.close()


@app.route('/api/banho/fila')
@login_required('banho')
def api_banho_fila():
    db = Session()
    try:
        fila = db.query(Card).filter_by(estado=ST_NA_FILA_BANHO).order_by(Card.id).all()
        ativos = db.query(Card).filter_by(estado=ST_EM_BANHO).order_by(Card.banho_inicio).all()
        return jsonify({
            'fila': [c.to_dict() for c in fila],
            'em_banho': [c.to_dict() for c in ativos],
        })
    finally:
        db.close()


@app.route('/api/banho/iniciar', methods=['POST'])
@login_required('banho')
def api_banho_iniciar():
    d = request.json or {}
    db = Session()
    try:
        card = db.query(Card).get(int(d.get('id', 0)))
        if not card or card.estado != ST_NA_FILA_BANHO:
            return jsonify({'sucesso': False, 'erro': 'Card não está na fila.'}), 404
        card.banho_inicio = datetime.utcnow()
        card.operador_banho = session.get('nome', '')
        card.estado = ST_EM_BANHO
        db.commit()
        return jsonify({'sucesso': True})
    finally:
        db.close()


@app.route('/api/banho/finalizar', methods=['POST'])
@login_required('banho')
def api_banho_finalizar():
    d = request.json or {}
    db = Session()
    try:
        card = db.query(Card).get(int(d.get('id', 0)))
        if not card or card.estado != ST_EM_BANHO:
            return jsonify({'sucesso': False, 'erro': 'Card não está em banho.'}), 404
        card.banho_fim = datetime.utcnow()
        card.banho_minutos = round((card.banho_fim - card.banho_inicio).total_seconds() / 60, 1)
        card.estado = ST_CONCLUIDO
        db.commit()
        return jsonify({'sucesso': True, 'banho_minutos': card.banho_minutos})
    finally:
        db.close()


@app.route('/api/dashboard/dados')
@login_required('admin')
def api_dashboard_dados():
    """Atualização em tempo real dos KPIs e do quadro de andamento."""
    db = Session()
    try:
        concluidos = db.query(Card).filter_by(estado=ST_CONCLUIDO).all()
        ativos = db.query(Card).filter(Card.estado != ST_CONCLUIDO)\
            .order_by(Card.id.desc()).all()
        return jsonify({
            'total': len(concluidos),
            'em_andamento': len(ativos),
            'ativos': [c.to_dict() for c in ativos],
        })
    finally:
        db.close()


# ─────────────────────────────────────────────────────────────────────────────
# Export Excel
# ─────────────────────────────────────────────────────────────────────────────
@app.route('/api/download_excel')
@login_required('admin')
def api_download_excel():
    db = Session()
    try:
        cards = db.query(Card).filter_by(estado=ST_CONCLUIDO).order_by(Card.id).all()
        wb = Workbook()
        ws = wb.active
        ws.title = 'Historico'
        cab = ['ID', 'Cesto', 'OP', 'Código', 'Descrição', 'Qtd', 'Processo', 'Tipo',
               'Operador Prep.', 'Prep. Início', 'Prep. Fim', 'Prep. (min)',
               'Líder', 'Op. Banho', 'Banho Início', 'Banho Fim', 'Banho (min)', 'Obs']
        fill_h = PatternFill("solid", fgColor="1F4E79")
        font_h = Font(bold=True, color="FFFFFF")
        for col, c in enumerate(cab, 1):
            cell = ws.cell(row=1, column=col, value=c)
            cell.fill = fill_h
            cell.font = font_h
            cell.alignment = Alignment(horizontal='center', vertical='center')
        for card in cards:
            dd = card.to_dict()
            ws.append([dd['id'], dd['numero_cesto'], dd['op'], dd['codigo'],
                       dd['descricao'], dd['quantidade'], dd['processo'], dd['tipo'],
                       dd['operador_prep'], dd['prep_inicio'], dd['prep_fim'],
                       dd['prep_minutos'], dd['lider'], dd['operador_banho'],
                       dd['banho_inicio'], dd['banho_fim'], dd['banho_minutos'],
                       dd['observacao']])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return send_file(buf, as_attachment=True,
                         download_name='historico_tratamento.xlsx',
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    finally:
        db.close()


@app.teardown_appcontext
def remove_session(exc=None):
    Session.remove()


init_db()

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    app.run(host='0.0.0.0', port=port, debug=True)
